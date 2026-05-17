from pathlib import Path
import csv
import json

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
ASSETS = OUT / "report_assets"
REPORT_PATH = ROOT / "HW2_report_chinese.docx"


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def fmt(value):
    try:
        return f"{float(value):.4f}"
    except Exception:
        return str(value)


def shade(cell, fill="D9EAF7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(header)
        shade(cell)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    return table


def add_picture(doc, path, width=5.6, caption=None):
    path = Path(path)
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def set_doc_fonts(doc):
    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[style_name]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    doc.styles["Normal"].font.size = Pt(10.5)


def main():
    task1 = read_csv(OUT / "task1" / "ablation_summary.csv")
    task3 = read_csv(OUT / "task3" / "loss_comparison.csv")
    pc_ce = read_csv(OUT / "task3" / "per_class_iou_ce.csv")
    pc_dice = read_csv(OUT / "task3" / "per_class_iou_dice.csv")
    pc_combined = read_csv(OUT / "task3" / "per_class_iou_combined.csv")
    tracking = json.load(open(OUT / "task2" / "tracking_summary.json", encoding="utf-8"))
    yolo_last = pd.read_csv(OUT / "task2" / "vehicle_finetune" / "results.csv").iloc[-1]

    doc = Document()
    set_doc_fonts(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("HW2 深度学习与空间智能实验报告")
    run.bold = True
    run.font.size = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("任务一：花卉分类 | 任务二：车辆检测跟踪与计数 | 任务三：U-Net 语义分割")

    doc.add_paragraph(
        "本报告根据 HW2_深度学习与空间智能.pdf 的要求，对三个实验任务的实验目标、"
        "实现方法、实验结果和结果分析进行整理。实验代码位于 code 目录，实验输出位于 outputs 目录。"
    )

    doc.add_heading("1. 实验环境与文件说明", level=1)
    add_table(doc, ["项目", "内容"], [
        ["深度学习框架", "PyTorch、torchvision、Ultralytics YOLOv8"],
        ["辅助库", "OpenCV、NumPy、Pillow、python-docx"],
        ["任务一数据集", "Oxford 102 Category Flower Dataset"],
        ["任务二数据集", "Road Vehicle Images Dataset 与交通视频 test_video.mp4"],
        ["任务三数据集", "Stanford Background Dataset，使用 labels_raw/*.regions.txt 作为 8 类语义标签"],
        ["结果目录", "outputs/task1、outputs/task2、outputs/task3"],
    ])

    doc.add_paragraph(
        "主要超参数：任务一训练 30 epochs，batch size 为 64，分类头学习率为 1e-3，"
        "backbone 学习率为 1e-4；任务二 YOLOv8 训练 50 epochs；任务三 U-Net 训练 40 epochs，"
        "batch size 为 8，学习率为 1e-3。"
    )

    doc.add_heading("2. 任务一：ImageNet 预训练 CNN 花卉分类", level=1)
    doc.add_paragraph(
        "任务要求是在 102 Category Flower Dataset 上建立 ResNet-18 或 ResNet-34 baseline，"
        "比较 ImageNet 预训练与从零训练，并加入 SE-block 或 CBAM 注意力模块进行性能对比。"
    )
    doc.add_heading("2.1 方法", level=2)
    doc.add_paragraph(
        "实验使用 ResNet-18 和 ResNet-34 作为主干网络，将最后分类头替换为 Dropout + Linear(102)。"
        "预训练模型加载 ImageNet 权重，随机初始化模型作为消融对照；SE 和 CBAM 插入到 ResNet 的 stage 后。"
        "训练中使用 AdamW、差异学习率、Label Smoothing、RandAugment、MixUp/CutMix、Warmup + Cosine 学习率调度。"
    )
    doc.add_heading("2.2 实验结果", level=2)
    add_table(doc, ["实验", "模型", "预训练", "注意力", "Best Val Acc", "Test Acc"], [
        [r["exp_name"], r["arch"], r["pretrained"], r["attention"], fmt(r["best_val_acc"]), fmt(r["test_acc"])]
        for r in task1
    ])
    add_picture(doc, ASSETS / "task1_accuracy.png", caption="图 1 任务一测试集准确率对比")
    add_picture(doc, ASSETS / "task1_val_curves.png", caption="图 2 任务一验证准确率曲线")
    doc.add_heading("2.3 分析", level=2)
    best_task1 = max(task1, key=lambda r: float(r["test_acc"]))
    cbam = [r for r in task1 if r["attention"] == "cbam"][0]
    doc.add_paragraph(
        f"预训练带来明显提升：ResNet-18 预训练测试准确率为 {fmt(task1[0]['test_acc'])}，"
        f"随机初始化仅为 {fmt(task1[1]['test_acc'])}。这说明 ImageNet 预训练特征对花卉小样本分类非常有效。"
    )
    doc.add_paragraph(
        f"注意力模块中，CBAM 的测试准确率为 {fmt(cbam['test_acc'])}，略高于普通 ResNet-18。"
        f"整体最优配置为 {best_task1['exp_name']}，测试准确率为 {fmt(best_task1['test_acc'])}。"
    )

    doc.add_heading("3. 任务二：YOLOv8 车辆检测、跟踪与跨线计数", level=1)
    doc.add_paragraph(
        "任务要求使用 Road Vehicle Images Dataset 训练 YOLOv8，随后对 10-30 秒车辆视频输出 Bounding Box "
        "和 Tracking ID，分析遮挡或 ID 变化帧，并统计跨越虚拟线的目标数量。"
    )
    doc.add_heading("3.1 方法", level=2)
    doc.add_paragraph(
        "检测部分使用 YOLOv8 微调车辆类别；视频阶段调用 YOLOv8 track 接口并使用 ByteTrack 维护 Tracking ID。"
        "跨线计数通过检测框中心点相对虚拟线的叉积符号变化判断是否过线，并用 crossed_ids 避免重复计数。"
        "遮挡事件通过框间 IoU 和 ID 重现间隔触发保存，并加入 event cooldown 避免连续帧重复保存相同事件。"
    )
    doc.add_heading("3.2 检测结果", level=2)
    add_table(doc, ["指标", "结果"], [
        ["Precision(B)", fmt(yolo_last["metrics/precision(B)"])],
        ["Recall(B)", fmt(yolo_last["metrics/recall(B)"])],
        ["mAP50(B)", fmt(yolo_last["metrics/mAP50(B)"])],
        ["mAP50-95(B)", fmt(yolo_last["metrics/mAP50-95(B)"])],
        ["训练 epoch", int(yolo_last["epoch"])],
    ])
    add_picture(doc, ASSETS / "task2_map_curve.png", caption="图 3 YOLOv8 验证集 mAP 曲线")
    add_picture(doc, OUT / "task2" / "vehicle_finetune" / "results.png", caption="图 4 YOLOv8 训练过程指标")
    add_picture(doc, OUT / "task2" / "vehicle_finetune" / "confusion_matrix_normalized.png", width=5.0, caption="图 5 归一化混淆矩阵")

    doc.add_heading("3.3 跟踪计数结果", level=2)
    add_table(doc, ["项目", "结果"], [
        ["输出视频", "outputs/task2/tracked_output.mp4"],
        ["虚拟线", f"{tracking['line_pt1']} -> {tracking['line_pt2']}"],
        ["跨线目标总数", tracking["total_cross_count"]],
        ["保存遮挡事件数", tracking["occlusion_events_saved"]],
        ["事件冷却间隔", f"{tracking.get('event_cooldown_frames', 0)} frames"],
        ["跨线 Tracking ID", ", ".join(map(str, tracking["crossed_ids"]))],
    ])
    add_picture(doc, OUT / "task2" / "occlusion_frames" / "event_03" / "frame_00090.jpg", width=4.8, caption="图 6 遮挡或密集车辆事件示例帧")

    doc.add_heading("3.4 视频分析步骤", level=2)
    video_steps = [
        "打开 outputs/task2/tracked_output.mp4，检查检测框、类别、置信度和 Tracking ID 是否显示完整。",
        "观察虚拟线附近车辆中心点轨迹，确认车辆跨线时 Count 增加，同一 ID 不重复计数。",
        "查看 outputs/task2/occlusion_frames/event_00 至 event_03 中的连续帧，分析车辆遮挡前后 ID 是否保持一致。",
        "若遮挡前后同一车辆 ID 不变，可说明 ByteTrack 保持了目标身份；若 ID 跳变，则记录为潜在 ID switch。",
        "结合 tracking_summary.json 汇总 total_cross_count、crossed_ids 和 occlusion_events_saved。",
    ]
    for step in video_steps:
        doc.add_paragraph(step)
    doc.add_heading("3.5 分析", level=2)
    doc.add_paragraph(
        f"检测模型最终 mAP50 为 {fmt(yolo_last['metrics/mAP50(B)'])}，"
        f"mAP50-95 为 {fmt(yolo_last['metrics/mAP50-95(B)'])}，能够完成车辆检测，"
        "但仍存在类别不均衡和小目标检测难点。"
        f"视频中最终统计到 {tracking['total_cross_count']} 个跨线目标，保存 "
        f"{tracking['occlusion_events_saved']} 组遮挡事件。"
    )

    doc.add_heading("4. 任务三：从零实现 U-Net 与 Loss 对比", level=1)
    doc.add_paragraph(
        "任务要求从零搭建 U-Net，包括下采样、上采样和 Skip Connection，在 Stanford Background Dataset 上训练，"
        "并比较 CE Loss、Dice Loss、CE + Dice Loss 的 mIoU。"
    )
    doc.add_heading("4.1 方法", level=2)
    doc.add_paragraph(
        "U-Net 由四层 Encoder、Bottleneck 和四层 Decoder 构成。Encoder 使用 ConvBlock + MaxPool，"
        "Decoder 使用反卷积上采样，并与对应 Encoder 特征进行 Skip Connection 拼接。输出层为 1x1 卷积，"
        "预测 8 类语义分割图。训练中使用同步水平翻转和 ColorJitter 增强，并输出每类 IoU。"
    )
    doc.add_heading("4.2 Loss 对比结果", level=2)
    add_table(doc, ["Loss", "Best Val mIoU"], [[r["loss_mode"], fmt(r["best_val_mIoU"])] for r in task3])
    add_picture(doc, ASSETS / "task3_loss_comparison.png", caption="图 7 不同 Loss 的 mIoU 对比")

    doc.add_heading("4.3 逐类 IoU", level=2)
    add_table(doc, ["Class ID", "Class Name", "CE", "Dice", "Combined"], [
        [row["class_id"], row["class_name"], fmt(row["iou"]), fmt(pc_dice[i]["iou"]), fmt(pc_combined[i]["iou"])]
        for i, row in enumerate(pc_ce)
    ])
    add_picture(doc, ASSETS / "task3_per_class_combined.png", caption="图 8 Combined Loss 下逐类 IoU")

    doc.add_heading("4.4 分析", level=2)
    best_task3 = max(task3, key=lambda r: float(r["best_val_mIoU"]))
    doc.add_paragraph(
        f"当前结果中，{best_task3['loss_mode']} 的 best val mIoU 最高，为 {fmt(best_task3['best_val_mIoU'])}。"
        "CE Loss 表现稳定，说明像素级分类监督对该数据集有效；Dice Loss 单独训练较低，"
        "可能因为多类语义分割中类别面积差异和边界复杂度导致优化不够稳定；Combined Loss 与 CE 接近，"
        "说明区域重叠约束有补充作用，但 CE/Dice 权重仍可继续搜索。"
    )
    doc.add_paragraph(
        "逐类结果显示 sky、road、building 等类别 IoU 较高，而 mountain 类 IoU 较低，"
        "说明少样本、小区域或外观变化较大的类别仍是主要难点。"
    )

    doc.add_heading("5. 总结与改进方向", level=1)
    doc.add_paragraph(
        "三个任务分别对应图像级分类、目标级检测跟踪和像素级分割。任务一验证了预训练与注意力机制的作用；"
        "任务二实现了车辆检测、跟踪 ID 和跨线计数；任务三完成了 U-Net 从零实现和损失函数对比。"
    )
    add_table(doc, ["任务", "已完成内容", "改进方向"], [
        ["任务一", "ResNet 预训练/随机初始化消融，SE/CBAM 对比，Accuracy 曲线", "尝试 ConvNeXt、EfficientNet、Swin Transformer，系统搜索增强和学习率"],
        ["任务二", "YOLOv8 微调、ByteTrack 跟踪、虚拟线计数、遮挡事件保存", "使用 YOLOv8s/m，提高小类别样本量，人工优化计数线位置"],
        ["任务三", "U-Net 从零实现，CE/Dice/Combined Loss 对比，逐类 IoU 分析", "尝试类别重加权、Attention U-Net、U-Net++、更多 CE/Dice 权重组合"],
    ])
    doc.add_paragraph("附：所有结果文件、训练曲线、输出视频、遮挡帧和统计表均保存在 outputs 目录。")

    doc.save(REPORT_PATH)
    print(REPORT_PATH)


if __name__ == "__main__":
    main()
